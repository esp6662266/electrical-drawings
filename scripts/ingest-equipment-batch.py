"""Read-only UNC intake, content-addressed originals, and static text inspection."""
from pathlib import Path
import json,hashlib,re,sys,zipfile,struct,zlib,xml.etree.ElementTree as ET
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'qa'/'hwpdeps'))
import olefile
from pypdf import PdfReader
SOURCE=Path(r'\\192.168.20.250\시설관리\통합지원본부\김범주(작업정리중)')
DOC=ROOT/'documents';STORE=DOC/'source-store';TEXT=DOC/'source-text';QA=ROOT/'qa'/'equipment-batch'
for d in [STORE,TEXT,QA]:d.mkdir(parents=True,exist_ok=True)
NAMES={3:'PREHEATER-03',4:'ALLOY50-04',5:'ALLOY30-05',6:'HOLDING-06',7:'ROLLING30-07',8:'ROLLING85-08',9:'LADLELIFT-09',10:'GBF-10',11:'CHIP-11',12:'DUST-12',13:'PRIMARY-13',14:'SHREDDER6060-14',15:'XRT-15'}
def digest(raw):return hashlib.sha256(raw).hexdigest()
def hwp_text(file):
    result=[]
    with olefile.OleFileIO(file) as ole:
        compressed=bool(struct.unpack_from('<I',ole.openstream('FileHeader').read(),36)[0]&1)
        for p in sorted(ole.listdir()):
            if p[0]!='BodyText':continue
            b=ole.openstream(p).read();b=zlib.decompress(b,-15) if compressed else b;pos=0
            while pos+4<=len(b):
                v=struct.unpack_from('<I',b,pos)[0];pos+=4;tag=v&1023;size=v>>20
                if size==4095:size=struct.unpack_from('<I',b,pos)[0];pos+=4
                payload=b[pos:pos+size];pos+=size
                if tag!=67:continue
                units=struct.unpack('<'+'H'*(len(payload)//2),payload[:len(payload)//2*2]);chars=[];i=0
                while i<len(units):
                    c=units[i]
                    if c>=32:chars.append(chr(c));i+=1
                    elif c in (10,13):chars.append('\n');i+=1
                    elif c==9:chars.append('\t');i+=8
                    elif c in (1,2,3,4,5,6,7,8,11,12,14,15,16,17,18,19,20,21,22,23):i+=8
                    else:i+=1
                text=''.join(chars).encode('utf-16le','surrogatepass').decode('utf-16le','replace').strip()
                if text:result.append(text)
    return result
def static_inspect(file,sha,ext):
    cache=TEXT/(sha+'.json')
    if cache.exists():
        try:return json.loads(cache.read_text(encoding='utf8'))
        except (ValueError,UnicodeError):pass
    out={'sha256':sha,'extension':ext,'status':'原型保管'.replace('原型保管','원형 보존'),'pages':[]}
    try:
        if ext=='.pdf':
            import pypdfium2 as pdfium
            reader=pdfium.PdfDocument(file);out['page_count']=len(reader)
            for i in range(len(reader)):
                page=reader[i]
                try:
                    textpage=page.get_textpage();text=textpage.get_text_range() or '';textpage.close()
                except Exception as ex:text='';out.setdefault('warnings',[]).append(f'PDF page {i+1}: {type(ex).__name__}')
                width,height=page.get_size();out['pages'].append({'page':i+1,'text':text,'width':width,'height':height});page.close()
            reader.close()
            out['status']='PDF 본문 추출; 시각 검토 별도'
        elif ext=='.hwp':
            out['paragraphs']=hwp_text(file);out['status']='HWP 본문 추출; 페이지 배치 미재현'
        elif ext in ('.docx','.pptx','.xlsx'):
            with zipfile.ZipFile(file) as z:
                if ext=='.pptx':
                    for name in sorted([n for n in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml$',n)],key=lambda n:int(re.search(r'(\d+)\.xml$',n)[1])):
                        tree=ET.fromstring(z.read(name));text='\n'.join(t.text or '' for t in tree.iter() if t.tag.endswith('}t'));out['pages'].append({'page':len(out['pages'])+1,'text':text})
                    out['page_count']=len(out['pages']);out['status']='슬라이드 텍스트 추출; 배치 미재현'
                elif ext=='.docx':
                    tree=ET.fromstring(z.read('word/document.xml'));out['paragraphs']=[''.join(t.text or '' for t in p.iter() if t.tag.endswith('}t')) for p in tree.iter() if p.tag.endswith('}p')];out['status']='DOCX 본문 추출; 페이지 배치 미재현'
                else:
                    import openpyxl
                    book=openpyxl.load_workbook(file,data_only=False,read_only=True,keep_links=False)
                    out['sheets']=[{'name':s.title,'rows':[[str(c) if c is not None else '' for c in row] for row in s.iter_rows(values_only=True) if any(c is not None for c in row)]} for s in book]
                    out['status']='XLSX 셀·수식 추출; 계산/매크로 실행 안 함';book.close()
        elif ext=='.xls':
            import xlrd
            book=xlrd.open_workbook(file,on_demand=True);out['sheets']=[{'name':s.name,'rows':[[str(v) for v in s.row_values(i)] for i in range(s.nrows) if any(str(v) for v in s.row_values(i))]} for s in book.sheets()];out['status']='XLS 저장 셀값 추출; 계산/매크로 실행 안 함';book.release_resources()
        elif ext=='.ppt':
            with olefile.OleFileIO(file) as ole:
                b=ole.openstream('PowerPoint Document').read();texts=[]
                for pos in range(0,len(b)-8):
                    _,typ,size=struct.unpack_from('<HHI',b,pos)
                    if typ in (4000,4008) and 0<size<200000 and pos+8+size<=len(b):
                        text=b[pos+8:pos+8+size].decode('utf-16le' if typ==4000 else 'cp949','replace').strip()
                        if len(text)>1 and sum(c.isprintable() or c.isspace() for c in text)/len(text)>.97:texts.append(text)
                out['paragraphs']=list(dict.fromkeys(texts));out['status']='PPT 텍스트 레코드 추출; 슬라이드 배치 미재현'
        elif ext in ('.zip','.xgwx'):
            if zipfile.is_zipfile(file):
                with zipfile.ZipFile(file) as z:out['archive_entries']=[{'name':i.filename,'bytes':i.file_size,'compressed_bytes':i.compress_size,'encrypted':bool(i.flag_bits&1)} for i in z.infolist()]
                out['status']='압축 목록만 정적 확인; 실행하지 않음'
            else:
                raw=file.read_bytes();out['header_hex']=raw[:64].hex();out['status']='원본 바이너리 보존; 전용 형식 정적 해석 필요'
        elif ext in ('.jpg','.jpeg','.png','.tif','.tiff'):
            from PIL import Image
            with Image.open(file) as im:out['image']={'width':im.width,'height':im.height,'frames':getattr(im,'n_frames',1)}
            out['status']='이미지 메타데이터 확인; 시각 검토 별도'
    except Exception as ex:out['error']=type(ex).__name__+': '+str(ex)[:250];out['status']='원형 보존; 자동 본문 추출 제한'
    cache.write_text(json.dumps(out,ensure_ascii=True,indent=2),encoding='utf8');return out
def main():
    catalog=[];unique={};summary=[]
    folders=sorted([p for p in SOURCE.iterdir() if p.is_dir() and re.match(r'^(\d+)번',p.name) and 3<=int(re.match(r'^(\d+)',p.name)[1])<=15],key=lambda p:int(re.match(r'^(\d+)',p.name)[1]))
    for folder in folders:
        number=int(re.match(r'^(\d+)',folder.name)[1]);items=[]
        for path in sorted(folder.rglob('*')):
            if not path.is_file():continue
            before=path.stat();raw=path.read_bytes();sha=digest(raw);ext=path.suffix.lower();dest=STORE/(sha+ext)
            if not dest.exists():dest.write_bytes(raw)
            assert digest(dest.read_bytes())==sha
            after=path.stat();assert before.st_size==after.st_size==len(raw) and before.st_mtime_ns==after.st_mtime_ns
            entry={'id':'S'+sha[:16],'name':path.name,'relative_path':path.relative_to(folder).as_posix(),'file':'/documents/source-store/'+dest.name,'sha256':sha,'bytes':len(raw),'extension':ext}
            if sha not in unique:unique[sha]=static_inspect(dest,sha,ext)
            info=unique[sha];entry['inspection']='/documents/source-text/'+sha+'.json';entry['inspection_status']=info['status']
            for key in ['page_count','error','image']:
                if key in info:entry[key]=info[key]
            items.append(entry)
        manifest={'equipment_id':NAMES[number],'number':number,'name':folder.name,'source_folder':str(folder),'source_count':len(items),'unique_source_count':len(set(s['sha256'] for s in items)),'sources':items}
        out=DOC/f'equipment-{number:02d}';out.mkdir(exist_ok=True);(out/'source-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8');catalog.append(manifest)
        summary.append({'number':number,'name':folder.name,'files':len(items),'unique':manifest['unique_source_count'],'pdf_pages':sum(unique[s]['page_count'] for s in set(i['sha256'] for i in items) if unique[s]['extension']=='.pdf' and 'page_count' in unique[s]),'extraction_errors':sum('error' in unique[s] for s in set(i['sha256'] for i in items))})
        print(json.dumps(summary[-1],ensure_ascii=True),flush=True)
    (DOC/'batch-catalog.json').write_text(json.dumps({'equipment':catalog,'summary':summary,'original_file_count':sum(c['source_count'] for c in catalog),'unique_hash_count':len(unique)},ensure_ascii=False,indent=2),encoding='utf8')
    print('INTAKE COMPLETE',sum(c['source_count'] for c in catalog),'references;',len(unique),'unique hashes',flush=True)

if __name__=="__main__":main()
